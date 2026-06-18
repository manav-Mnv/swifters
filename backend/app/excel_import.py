from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

REQUIRED_HEADERS = {
    "enrollment number": "enrollment",
    "full name": "name",
    "department": "department",
    "course": "course",
    "institute": "institute",
}

OPTIONAL_HEADERS = {
    "email": "email",
    "phone": "phone",
}


def parse_members_from_workbook(file_bytes: bytes) -> tuple[list[dict], list[str]]:
    workbook = load_workbook(filename=BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Workbook is empty")

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    normalized = {header.lower(): index for index, header in enumerate(headers) if header}
    missing = [header for header in REQUIRED_HEADERS if header not in normalized]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")

    seen: set[str] = set()
    members: list[dict] = []
    warnings: list[str] = []
    column_map = {**REQUIRED_HEADERS, **OPTIONAL_HEADERS}

    for row_number, row in enumerate(rows[1:], start=2):
        if not row or all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        member = {}
        for source_header, target_key in column_map.items():
            index = normalized.get(source_header)
            value = row[index] if index is not None and index < len(row) else None
            member[target_key] = "" if value is None else str(value).strip()
        enrollment = member["enrollment"]
        if not enrollment:
            warnings.append(f"Row {row_number} skipped: missing enrollment number")
            continue
        if enrollment in seen:
            warnings.append(f"Row {row_number} skipped: duplicate enrollment {enrollment} in file")
            continue
        seen.add(enrollment)
        members.append(member)

    return members, warnings
